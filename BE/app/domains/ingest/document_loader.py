"""
Tải và chunk tài liệu bằng LangChain — bổ sung (không thay) extract/split legacy trong ingest_utils.
Video / OCR vẫn dùng ingest_utils.extract_text khi cần.
"""

from __future__ import annotations

import json
import os
from pathlib import Path
from typing import List

from langchain_core.documents import Document

from app.domains.ingest import formats

# Giữ tên cũ cho chỗ nào còn tham chiếu, nhưng nội dung lấy từ formats.py.
SUPPORTED_LC_EXTENSIONS = formats.SUPPORTED_EXTENSIONS


def _single_doc(text: str, source_stem: str, file_path: str) -> List[Document]:
    t = (text or "").strip()
    if not t:
        return []
    return [Document(page_content=t, metadata={"source": source_stem, "file_path": file_path})]


def extract_docx_blocks(path: Path) -> list[dict]:
    """Extract non-empty DOCX paragraphs and table rows in document order.

    Stable paragraph/table indices are retained so reviewed source-unit locators can
    be bridged to the exact text used by the production loader.
    """
    from docx import Document as DocxDocument
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = DocxDocument(str(path))
    blocks: list[dict] = []
    paragraph_index = 0
    table_index = 0
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, doc)
            text = (paragraph.text or "").strip()
            if text:
                blocks.append({
                    "locator": {
                        "type": "docx_paragraph",
                        "paragraph_index": paragraph_index,
                    },
                    "text": text,
                })
            paragraph_index += 1
        elif isinstance(child, CT_Tbl):
            table = Table(child, doc)
            for row_index, row in enumerate(table.rows):
                text = " | ".join((cell.text or "").strip() for cell in row.cells).strip()
                if text:
                    blocks.append({
                        "locator": {
                            "type": "docx_table_row",
                            "table_index": table_index,
                            "row_index": row_index,
                        },
                        "text": text,
                    })
            table_index += 1
    return blocks


def _load_docx_python(path: Path) -> List[Document]:
    blocks = extract_docx_blocks(path)
    text = "\n".join(block["text"] for block in blocks)
    return _single_doc(text, path.stem, str(path))


def _load_pymupdf(path: Path) -> List[Document]:
    """Đọc theo trang bằng PyMuPDF. Ngoài PDF, fitz mở native cả EPUB/MOBI/FB2/XPS
    nên dùng chung một vòng lặp — không cần nhánh riêng cho từng định dạng sách."""
    import fitz

    documents: List[Document] = []
    with fitz.open(path) as pdf:
        for page_index, page in enumerate(pdf):
            text = page.get_text() or ""
            if text.strip():
                documents.append(Document(
                    page_content=text,
                    metadata={"source": path.stem, "file_path": str(path), "page": page_index},
                ))
    return documents


def _soffice_convert(path: Path, target_ext: str, out_dir: str) -> Path:
    """Nhờ LibreOffice đổi định dạng lạ sang OOXML rồi đọc bằng thư viện Python.

    Dùng cho .odt/.odp (và .doc ở markdown_convert). Không có `soffice` trong
    PATH thì file đích không sinh ra và hàm ném FileNotFoundError — caller bắt
    rồi rơi về extract_text như mọi nhánh lỗi khác.
    """
    import subprocess

    subprocess.run(
        ["soffice", "--headless", "--convert-to", target_ext, str(path), "--outdir", out_dir],
        check=False,
        capture_output=True,
        timeout=180,
    )
    converted = Path(out_dir) / f"{path.stem}.{target_ext}"
    if not converted.exists():
        raise FileNotFoundError(f"soffice khong doi duoc {path.name} sang .{target_ext}")
    return converted


def _load_html(path: Path) -> List[Document]:
    """HTML sang Markdown bằng markdownify.

    Bản cũ gọi `UnstructuredHTMLLoader` mà `unstructured` KHÔNG có trong
    requirements.txt — lỗi import bị `except Exception: pass` nuốt, rơi xuống
    extract_text (không biết html) nên mọi file .html trả rỗng và ingest chết
    với "Cannot read file content". markdownify đã cài sẵn và markdown_convert
    vẫn đang dùng nó cho đúng việc này.
    """
    from markdownify import markdownify as html_to_markdown

    raw = path.read_text(encoding="utf-8", errors="ignore")
    return _single_doc(html_to_markdown(raw), path.stem, str(path))


def _load_pptx(path: Path) -> List[Document]:
    """Mỗi slide một Document. Lấy text trong shape, ô bảng, và ghi chú người trình bày.

    Ghi chú slide thường chứa phần giảng giải không có trên slide nên giữ lại,
    đánh dấu rõ để về sau truy vết được nguồn.
    """
    from pptx import Presentation

    documents: List[Document] = []
    deck = Presentation(str(path))
    for index, slide in enumerate(deck.slides):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = (shape.text_frame.text or "").strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    row_text = " | ".join((c.text or "").strip() for c in row.cells).strip()
                    if row_text:
                        parts.append(row_text)
        if slide.has_notes_slide:
            note = (slide.notes_slide.notes_text_frame.text or "").strip()
            if note:
                parts.append(f"[Ghi chú slide] {note}")
        body = "\n".join(parts).strip()
        if body:
            documents.append(Document(
                page_content=body,
                metadata={"source": path.stem, "file_path": str(path), "slide": index + 1},
            ))
    return documents


def _load_xlsx(path: Path) -> List[Document]:
    """Mỗi sheet một Document, nội dung là bảng Markdown.

    Bảng Markdown thay vì CSV thô vì bước Normalize/chunk sau đó cắt theo cấu
    trúc Markdown — giữ cùng một ngôn ngữ bảng thì heading và hàng không bị vỡ.
    Đọc `read_only` + `data_only` để lấy giá trị đã tính, không phải công thức.
    """
    from openpyxl import load_workbook

    documents: List[Document] = []
    book = load_workbook(str(path), read_only=True, data_only=True)
    try:
        for sheet in book.worksheets:
            rows = [
                [("" if c is None else str(c)).strip() for c in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            rows = [r for r in rows if any(r)]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            header, body = rows[0], rows[1:]
            lines = ["| " + " | ".join(header) + " |", "|" + "---|" * width]
            lines += ["| " + " | ".join(r) + " |" for r in body]
            documents.append(Document(
                page_content=f"## {sheet.title}\n\n" + "\n".join(lines),
                metadata={"source": path.stem, "file_path": str(path), "sheet": sheet.title},
            ))
    finally:
        book.close()
    return documents


def _load_rtf(path: Path) -> List[Document]:
    from striprtf.striprtf import rtf_to_text

    raw = path.read_text(encoding="utf-8", errors="ignore")
    return _single_doc(rtf_to_text(raw, errors="ignore"), path.stem, str(path))


def _load_opendocument(path: Path) -> List[Document]:
    """.odt/.odp: nhờ soffice đổi sang .docx/.pptx rồi dùng lại loader có sẵn."""
    import tempfile

    target = "pptx" if path.suffix.lower() == ".odp" else "docx"
    with tempfile.TemporaryDirectory() as temp_dir:
        converted = _soffice_convert(path, target, temp_dir)
        docs = _load_pptx(converted) if target == "pptx" else _load_docx_python(converted)
    # Metadata đang trỏ vào file tạm đã bị xoá — trả về đường dẫn thật.
    for doc in docs:
        doc.metadata["source"] = path.stem
        doc.metadata["file_path"] = str(path)
    return docs


def load_document(file_path: str) -> List[Document]:
    """
    Load file thành list Document. Luôn gắn metadata source (stem) + file_path.
    Fallback về ingest_utils.extract_text nếu loader LangChain thất bại.
    """
    path = Path(file_path)
    if not path.is_file():
        return []
    ext = path.suffix.lower()
    stem = path.stem
    fp = str(path)

    # OCR ảnh và Word 97-2003 đi thẳng qua extract_text (Tesseract / soffice).
    if ext in formats.IMAGE or ext == ".doc":
        from app.domains.ingest.ingest_utils import extract_text
        return _single_doc(extract_text(fp), stem, fp)

    try:
        if ext in formats.PYMUPDF:
            return _load_pymupdf(path)

        if ext in formats.TEXT:
            from langchain_community.document_loaders import TextLoader

            return TextLoader(fp, encoding="utf-8", autodetect_encoding=True).load()

        if ext == ".docx":
            return _load_docx_python(path)

        if ext in formats.WEB:
            return _load_html(path)

        if ext in formats.PPTX:
            return _load_pptx(path)

        if ext in formats.XLSX:
            return _load_xlsx(path)

        if ext in formats.RICH_TEXT:
            return _load_rtf(path)

        if ext in formats.OPENDOCUMENT:
            return _load_opendocument(path)

        if ext == ".csv":
            from langchain_community.document_loaders import CSVLoader

            return CSVLoader(fp).load()

        if ext == ".json":
            raw = path.read_text(encoding="utf-8", errors="ignore")
            try:
                obj = json.loads(raw)
                text = json.dumps(obj, ensure_ascii=False, indent=2) if not isinstance(obj, str) else obj
            except Exception:
                text = raw
            return _single_doc(text, stem, fp)

    except Exception:
        pass

    from app.domains.ingest.ingest_utils import extract_text
    return _single_doc(extract_text(fp), stem, fp)


def split_documents(
    docs: List[Document],
    chunk_size: int | None = None,
    chunk_overlap: int | None = None,
) -> List[Document]:
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    cs = chunk_size if chunk_size is not None else int(os.getenv("CHUNK_SIZE", "500"))
    co = chunk_overlap if chunk_overlap is not None else int(os.getenv("CHUNK_OVERLAP", "50"))
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=cs,
        chunk_overlap=co,
        separators=["\n\n", "\n", ".", "!", "?", ",", " ", ""],
        length_function=len,
    )
    out = splitter.split_documents(docs)
    for i, d in enumerate(out):
        d.metadata.setdefault("chunk_index", i)
        d.metadata["total_chunks"] = len(out)
    return out
