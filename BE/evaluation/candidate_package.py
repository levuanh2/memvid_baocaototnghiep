from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path


ROOT = Path("reports/evaluation/datasets/corpus_v1")


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def source_units(root: Path) -> list[dict]:
    import fitz
    from docx import Document

    documents = [json.loads(x) for x in (root / "documents.jsonl").read_text(encoding="utf-8").splitlines() if x.strip()]
    units: list[dict] = []
    for record in documents:
        if not record.get("eligible_for_study"):
            continue
        source = Path(record["source"])
        if record["format"] == "pdf":
            with fitz.open(source) as pdf:
                for page_number, page in enumerate(pdf, 1):
                    text = norm(page.get_text())
                    units.append({
                        "doc_id": record["doc_id"], "source": str(source),
                        "locator": {"type": "pdf_page", "page": page_number},
                        "locator_label": f"page {page_number}", "text": text,
                        "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
                    })
        elif record["format"] == "docx":
            doc = Document(source)
            heading = ""
            for index, paragraph in enumerate(doc.paragraphs):
                text = norm(paragraph.text)
                if not text:
                    continue
                if paragraph.style.name.lower().startswith("heading"):
                    heading = text
                units.append({
                    "doc_id": record["doc_id"], "source": str(source),
                    "locator": {"type": "docx_paragraph", "paragraph_index": index},
                    "locator_label": f"paragraph {index}; heading: {heading or '(document title)'}",
                    "heading": heading, "text": text,
                    "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
                })
            for table_index, table in enumerate(doc.tables):
                for row_index, row in enumerate(table.rows):
                    text = norm(" | ".join(norm(cell.text) for cell in row.cells))
                    units.append({
                        "doc_id": record["doc_id"], "source": str(source),
                        "locator": {"type": "docx_table_row", "table_index": table_index, "row_index": row_index},
                        "locator_label": f"table {table_index}, row {row_index}", "text": text,
                        "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
                    })
    return units


def e(kind: str, number: int, needle: str) -> tuple[str, int, str]:
    return kind, number, norm(needle)


def q(qid: str, doc: str, qtype: str, question: str, answer: str, evidence: list[tuple[str, int, str]],
      rationale: str, *, difficulty: str = "medium", answerability: str = "answerable",
      multi: bool = False, hierarchical: bool = False, citation: bool = False,
      contradiction: bool = False, ambiguity_reason: str = "", interpretations: list[str] | None = None,
      resolution: str = "", nearest: str = "") -> dict:
    return locals()


def definitions() -> list[dict]:
    V, R, S = "cv1_vector_feature_store", "cv1_ragas_guardrails", "cv1_social_engineering"
    rows = [
        q("V01",V,"exact-term factual","Khi đổi embedding model, hệ thống phải làm gì với chỉ mục?","Phải tái lập chỉ mục toàn bộ và dùng cùng phiên bản model ở lúc lập chỉ mục và lúc truy vấn.",[e("page",6,"Đổi embedding model → phải reindex toàn bộ. Train (index time) và serve (query time) phải dùng cùng model version")],"Kiểm tra một quy tắc vận hành được nêu trực tiếp.",citation=True),
        q("V02",V,"exact-term factual","Khi embedding đã unit-normalized, cosine, dot product và Euclidean cho thứ hạng như thế nào?","Ba độ đo cho cùng thứ tự xếp hạng.",[e("page",7,"unit-norm -> thứ tự xếp hạng (ranking) của cosine similarity, dot product và Euclidean distance là giống nhau")],"Kiểm tra chi tiết kỹ thuật có điều kiện."),
        q("V03",V,"exact-term factual","Hybrid Search trên slide kết hợp ba thành phần nào?","BM25, vector search và RRF.",[e("page",13,"Hybrid Search: BM25 + Vector + RRF")],"Câu hỏi thuật ngữ chính xác."),
        q("V04",V,"exact-term factual","Ví dụ RAG end-to-end dùng kích thước chunk và overlap bao nhiêu?","512 tokens và overlap 50.",[e("page",14,"512 tokens 50 overlap")],"Kiểm tra giá trị cấu hình trên sơ đồ; cần người xem slide xác nhận ngữ cảnh."),
        q("V05",V,"paraphrase factual","Vì sao việc chọn mô hình embedding ban đầu cần thận trọng?","Vì đổi mô hình embedding làm phát sinh chi phí tái lập chỉ mục cao.",[e("page",5,"Re-index cost cao thận từ đầu. ⇒ chọn cẩn")],"Diễn đạt lại hệ quả vận hành từ slide."),
        q("V06",V,"paraphrase factual","Feature store phục vụ dữ liệu huấn luyện và dữ liệu online theo hai loại kho nào?","Offline store và online store.",[e("page",22,"Đặc điểm của Online và Offline Store")],"Dùng nội dung bài tập như gợi ý thuật ngữ; người duyệt cần kiểm tra slide hình ảnh liên quan."),
        q("V07",V,"multi-chunk","Mô tả chuỗi truy hồi hai giai đoạn trong RAG và vị trí của reranking.","Truy vấn được embedding để tìm top-k trong vector DB, sau đó các ứng viên được rerank trước khi xây context cho LLM.",[e("page",14,"User Query Documents Query Embed Chunking Retrieve Top-K Embedding Vector DB LLM Generate"),e("page",16,"2-Stage Retrieval: Vector Search + Reranking")],"Cần kết hợp sơ đồ pipeline và slide reranking.",multi=True,hierarchical=True,citation=True),
        q("V08",V,"multi-chunk","Feature store xử lý khác nhau thế nào giữa dữ liệu huấn luyện và phục vụ trực tuyến?","Ứng viên trả lời cần đối chiếu offline/online store và point-in-time join; chi tiết phải được người duyệt đọc trực tiếp trên slide.",[e("page",21,"Feature Store Architecture"),e("page",23,"Online vs Offline Store + Point-in-Time Join")],"Nội dung phần lớn nằm trong sơ đồ; đây là ứng viên khó cần xác minh bằng hình ảnh.",difficulty="hard",multi=True,hierarchical=True),
        q("V09",V,"cross-section","Khi nào GraphRAG phù hợp hơn truy hồi đoạn văn thông thường, và nó được xây từ gì?","GraphRAG phù hợp khi quan hệ quan trọng hơn từng đoạn riêng lẻ và được xây từ document knowledge graph.",[e("page",17,"GraphRAG: Khi Quan Hệ Quan Trọng Hơn Đoạn Văn"),e("page",18,"GraphRAG Construction: Document Knowledge Graph")],"Liên kết hai slide kế tiếp về động cơ và cấu trúc.",multi=True,hierarchical=True),
        q("V10",V,"overview","Tài liệu Vector Store & Feature Store bao quát những chủ đề chính nào?","Embeddings, vector databases, ANN/hybrid retrieval, RAG, GraphRAG/knowledge graphs và feature stores.",[e("page",2,"Vector Embeddings: Text -> Số Vector DB Landscape ANN & Retrieval Patterns (HNSW, Hybrid) RAG Pipeline với Vector Store GraphRAG & Knowledge Graphs Feature Store")],"Câu hỏi tổng quan dựa trên mục lục.",hierarchical=True),
        q("V11",V,"main-points","Những điểm chính người học cần ghi nhớ ở cuối bài là gì?","Ứng viên cần được người duyệt trích từ slide Key Takeaways vì phần lớn nội dung có thể nằm trong đồ họa.",[e("page",26,"Tổng kết — Key Takeaways")],"Không suy diễn nội dung không trích xuất được; giữ làm ứng viên kiểm tra trực quan.",difficulty="hard",hierarchical=True),
        q("V12",V,"compare","Vector store và feature store được đặt trong hai kỷ nguyên ML/LLM như thế nào?","Ứng viên cần đối chiếu trực tiếp bảng ở slide 'ML Era vs LLM Era'.",[e("page",25,"Vector & Feature Store: ML Era vs LLM Era")],"Bảng không trích xuất đủ chữ; bắt buộc duyệt trực quan.",difficulty="hard",hierarchical=True),
        q("V13",V,"how/why","Vì sao training-serving skew được xem là một lỗi thầm lặng?","Ứng viên trả lời cần giải thích sự lệch giữa feature lúc huấn luyện và lúc phục vụ; slide chỉ cho tiêu đề nên cần người duyệt xác minh.",[e("page",24,"Training-Serving Skew: Lỗi Thầm Lặng")],"Câu hỏi nhân quả, không biến suy đoán thành nhãn.",difficulty="hard"),
        q("V14",V,"ambiguous","Store nào nên được dùng cho hệ thống này?","",[],"Từ 'store' có thể chỉ vector store, online feature store hoặc offline feature store.",answerability="ambiguous",ambiguity_reason="Thuật ngữ 'store' và use case đều chưa xác định.",interpretations=["vector store cho truy hồi embedding","online feature store cho serving","offline feature store cho training"],resolution="Cần làm rõ tác vụ, độ trễ và loại dữ liệu."),
        q("V15",V,"insufficient evidence","Giá trị M tối ưu của HNSW cho corpus tiếng Việt này là bao nhiêu?","",[],"Tài liệu yêu cầu tìm hiểu HNSW nhưng không nêu giá trị M tối ưu cho corpus cụ thể.",answerability="insufficient_evidence",difficulty="hard",nearest="Trang 11 chỉ nêu bài tập tìm hiểu HNSW, IVF, PQ, OPQ, IVF-PQ."),
        q("V16",V,"contradictory evidence","Slide có mâu thuẫn khi vừa cảnh báo chọn sai similarity metric làm giảm recall, vừa nói ba metric có cùng ranking không?","Không nhất thiết; kết luận cùng ranking chỉ áp dụng cho vector unit-normalized, còn cảnh báo metric sai có phạm vi rộng hơn.",[e("page",7,"Match metric với pretraining ● objective: sai metric degrade ● recall 10–20%"),e("page",7,"unit-norm -> thứ tự xếp hạng (ranking) của cosine similarity, dot product và Euclidean distance là giống nhau")],"Ứng viên conditional-difference cho NLI, không gán contradiction thật trước human review.",difficulty="hard",multi=True,contradiction=True),
        q("V17",V,"citation-sensitive","Cấu hình minh họa của pipeline RAG gồm model embedding và metadata filter nào?","Slide nêu text-embed-3/bge-m3 và lọc metadata theo source, date từ 2024.",[e("page",14,"text-embed-3 bge-m3 Metadata filter: source, date ≥ 2024")],"Các tên model/điều kiện cần citation chính xác.",citation=True),
        q("V18",V,"citation-sensitive","Tài liệu quy định gì về phiên bản embedding giữa index time và query time?","Hai thời điểm phải dùng cùng phiên bản mô hình.",[e("page",6,"Train (index time) và serve (query time) phải dùng cùng model version")],"Khẳng định vận hành cần trích nguồn trực tiếp.",citation=True),

        q("R01",R,"exact-term factual","Bốn metric cốt lõi của RAGAS trên slide là gì?","Faithfulness, Answer Relevancy, Context Precision và Context Recall.",[e("page",16,"Faithfulness Answer ↔Context (hallucination) Answer Relevancy Answer ↔Question (on-topic) Context Precision Retrieved chunks ranked (NDCG) Context Recall Coverage with ground truth (completeness)")],"Liệt kê trực tiếp bốn metric."),
        q("R02",R,"exact-term factual","Faithfulness được tính theo công thức nào?","Số claim được xác minh đúng chia cho tổng số claim.",[e("page",17,"Score = (verified True) / (total claims)")],"Câu hỏi công thức trực tiếp."),
        q("R03",R,"exact-term factual","Answer Relevancy dùng bao nhiêu reverse questions trong thuật toán minh họa?","Ba reverse questions.",[e("page",18,"Tạo n = 3 reverse questions từA")],"Kiểm tra tham số của thuật toán minh họa."),
        q("R04",R,"exact-term factual","Context Precision nhấn mạnh điều gì ngoài precision đơn thuần?","Nó đánh giá thứ hạng theo NDCG, ưu tiên chunk liên quan ở top.",[e("page",19,"Không chỉlà precision đơn thuần. Là NDCG (Normalized Discounted Cumulative Gain) — relevant chunks phải ởtop.")],"Kiểm tra bản chất ranking của metric."),
        q("R05",R,"exact-term factual","Ngưỡng Cohen's kappa nào được slide coi là production minimum?","Từ 0,60 đến 0,80 được gọi là substantial và là production minimum.",[e("page",34,"0.60 −0.80 Substantial Production minimum")],"Giá trị định lượng cần citation."),
        q("R06",R,"exact-term factual","Tổng latency user-facing của kiến trúc guardrail được phân bổ tối đa bao nhiêu?","Không quá 80 ms theo slide latency budget.",[e("page",45,"Total user-facing ≤80ms")],"Giá trị định lượng trực tiếp.",citation=True),
        q("R07",R,"paraphrase factual","Vì sao golden set vẫn cần thiết khi đã có metric reference-free?","Reference-free scale tốt nhưng có thể không bắt lỗi đúng theo ngữ cảnh; golden set reference-based được dùng cho regression.",[e("page",9,"Reference-free Không cần ground truth."),e("page",9,"Dùng cảhai. Reference-based cho golden set 100–500 q (regression test). Reference-free cho production sampling 5%.")],"Diễn giải lý do kết hợp hai loại eval.",multi=True),
        q("R08",R,"paraphrase factual","Offline và online evaluation bổ sung cho nhau như thế nào?","Offline dùng dataset cố định trước deploy/ở mỗi PR để chặn regression; online lấy traffic sau deploy để theo dõi drift và hành vi thực.",[e("page",10,"Offline eval When: trước deploy, mỗi PR. Where: dataset cốđịnh. Why: CI gate + regression detection."),e("page",10,"Online eval When: sau deploy, continuous. Where: sample 1–5% production traffic. Why: drift detection + monitoring.")],"Đối chiếu hai chế độ trên cùng slide.",multi=True),
        q("R09",R,"paraphrase factual","Cách giảm thiên lệch vị trí khi judge so sánh hai câu trả lời là gì?","Đánh giá cả (A,B) và (B,A) rồi lấy trung bình; hoặc ngẫu nhiên hóa thứ tự và cho phép tie.",[e("page",30,"Swap-and-average: eval cả(A,B) và (B,A), average score."),e("page",30,"Random ordering: mỗi eval call randomize."),e("page",30,"Tie option: cho phép judge trả“tie” khi unsure.")],"Paraphrase các biện pháp giảm position bias.",citation=True),
        q("R10",R,"paraphrase factual","Tại sao chỉ chặn output ở lượt tấn công chưa đủ để chống session poisoning?","Vì input độc hại vẫn nằm trong history và có thể được thực thi ở lượt sau; cần thay thế nó ngay trong lịch sử.",[e("page",53,"Block ởTurn 2 chỉblock output. Input đã vào history. Turn 3 agent treats history as trusted context →obey malicious request."),e("page",54,"Solution: Input-level replacement, không chỉoutput blocking.")],"Liên kết mô tả tấn công và biện pháp phòng thủ.",multi=True),
        q("R11",R,"multi-chunk","Tại sao đánh giá final answer chưa đủ cho agent, và cần bổ sung những metric nào?","Agent có thể trả lời đúng nhưng chọn sai tool, tốn chi phí hoặc làm lộ query; cần trajectory correctness, tool selection accuracy, step efficiency, cost per task và final answer quality.",[e("page",13,"Cuối cùng trảlời đúng ■Tốn $0.50 (thay vì $0.005), lộquery qua public Google"),e("page",13,"Trajectory correctness ■Tool selection accuracy ■Step efficiency ■Cost per task ■Final answer quality")],"Cần kết hợp ví dụ và danh sách metric.",multi=True),
        q("R12",R,"multi-chunk","So sánh NLI và SelfCheckGPT trong phát hiện hallucination.","NLI kiểm tra entailment/contradiction/neutral giữa context và câu trả lời; SelfCheckGPT lấy nhiều mẫu để đo tính nhất quán, không cần ground truth nhưng tốn nhiều lần gọi hơn.",[e("page",38,"Pattern thông minh không cần ground truth."),e("page",38,"Cost: 6x normal (1 + 5 samples)"),e("page",39,"premise + hypothesis →entailment / contradiction / neutral")],"So sánh hai phương pháp ở các phần khác nhau.",difficulty="hard",multi=True),
        q("R13",R,"multi-chunk","Kiến trúc defense-in-depth phân guardrail thành những trục và lớp nào?","Bốn trục là topical, safety, security, compliance; bốn lớp là input, LLM, output và audit.",[e("page",43,"Guardrails 4 Trục"),e("page",44,"L1 — Input Layer (< 30ms) L2 — LLM Layer (system prompt rules, 0ms) L3 — Output Layer (< 50ms) L4 — Audit Layer (async, không block)")],"Kết hợp taxonomy chức năng và kiến trúc triển khai.",multi=True,hierarchical=True),
        q("R14",R,"cross-section","Deliverable cuối ngày liên hệ thế nào với eval gate trong CI/CD?","Deliverable yêu cầu eval suite, judge, guardrails và blueprint; CI/CD biến chúng thành các bước smoke, RAGAS, judge và red-team có tiêu chí pass/fail.",[e("page",5,"Eval suite (RAGAS ≥0.75) + guardrail layer (overhead < 100ms P95) + blueprint document."),e("page",61,"Pattern: mọi PR chạy eval trước khi merge.")],"Liên kết mục tiêu học tập với quy trình production.",multi=True,hierarchical=True),
        q("R15",R,"cross-section","Tại sao test set cố định vẫn phải được cập nhật bằng lỗi production?","Test set có thể cũ và không phản ánh usage hiện tại; vòng Failure → test case bổ sung lỗi production vào regression suite để ngăn tái phát.",[e("page",23,"Test set staleness: dataset 6 tháng tuổi không reflect current usage."),e("page",62,"Pattern: Failure →test case loop")],"Kết hợp pitfall và quy trình cải tiến.",multi=True,hierarchical=True),
        q("R16",R,"overview","Tài liệu RAGAS & Guardrails bao quát các nhóm nội dung nào?","Foundations of evaluation, RAGAS, LLM-as-Judge và bias, hallucination detection, guardrails, prompt injection/output protection, production patterns và lab blueprint.",[e("page",3,"1. Foundations of Evaluation 2. RAGAS Deep Dive (4 core metrics) 3. LLM-as-Judge & 4 biases 4. Hallucination Detection 5. Guardrails Foundations 6. Prompt Injection & Output Guardrails 7. Production Patterns (CI/CD, compliance) 8. Lab 24: Eval + Guardrail blueprint")],"Tổng quan dựa trên mục lục.",hierarchical=True),
        q("R17",R,"main-points","Ba key takeaway cuối bài là gì?","Evaluation là bắt buộc với RAGAS và LLM-Judge; guardrails cần defense-in-depth; LLM-Judge có bốn bias cần cross-judge và hiệu chỉnh với human.",[e("page",71,"Eval ̸= optional. RAGAS 4 metrics + LLM-Judge là baseline."),e("page",71,"Defense-in-depth. Guardrails 4 layers"),e("page",71,"LLM-Judge có 4 biases. Position, length, self-enhancement, style")],"Tóm tắt đúng ba takeaway được đánh số.",multi=True,hierarchical=True),
        q("R18",R,"compare","Pairwise và absolute scoring khác nhau thế nào?","Absolute chấm một câu trả lời trên thang rubric, dễ so sánh qua các run nhưng chủ quan; pairwise chọn A/B/tie, ổn định hơn cho regression nhưng cần baseline.",[e("page",29,"Absolute scoring Score 1 answer trên rubric (1–5 scale)."),e("page",29,"Pairwise comparison Compare A vs B, pick winner (hoặc tie).")],"Đối chiếu hai phương pháp judge.",multi=True),
        q("R19",R,"compare","Khi nào nên chọn RAGAS, DeepEval hoặc Phoenix theo bảng so sánh?","RAGAS cho dự án tập trung RAG; DeepEval cho workflow Python/pytest; Phoenix khi cần observability và tracing production.",[e("page",25,"RAGAS là default — ecosystem mature, doc tốt, framework-agnostic. DeepEval nếu team đã dùng pytest. Phoenix nếu integrate với Day 13 observability stack.")],"So sánh công cụ theo use case."),
        q("R20",R,"compare","Llama Guard 3 và GCP Model Armor khác nhau về triển khai, latency và đánh đổi nào?","Llama Guard 3 là classifier OSS 8B tự host khoảng 40 ms trên A100; Model Armor là dịch vụ GCP managed khoảng 50–100 ms, có SLA/audit nhưng tốn phí và lock-in.",[e("page",55,"8B params, runs trên 1 GPU ■Latency ∼40ms (A100)"),e("page",57,"Pricing: $0.001–0.005/check ■Latency: 50–100ms (network) ■Vendor lock-in (GCP only)")],"So sánh hai lựa chọn output/enterprise guardrail.",difficulty="hard",multi=True,citation=True),
        q("R21",R,"how/why","Vì sao thiết kế eval là một bài toán trade-off giữa bias, variance và chi phí?","Mở rộng test set giúp giảm variance nhưng tăng chi phí, còn tăng human review giúp giảm bias nhưng tăng latency.",[e("page",14,"Tăng test size →giảm variance nhưng tăng cost. Tăng human review →giảm bias nhưng tăng latency.")],"Câu hỏi nhân quả trực tiếp."),
        q("R22",R,"how/why","Vì sao chuỗi input guardrail nên redact PII trước khi kiểm tra prompt injection?","Vì prompt injection có thể chứa PII; redact trước ngăn dữ liệu nhạy cảm tiếp tục đi qua validator sau.",[e("page",46,"Order matters: PII redact trước injection check (injection có thểchứa PII)")],"Lý do thứ tự validator được nêu rõ."),
        q("R23",R,"how/why","Tại sao guardrail quá chặt có thể làm hệ thống kém an toàn về lâu dài?","False positive làm người dùng khó chịu, học cách rephrase để bypass hoặc bỏ hệ thống; cần đo refuse rate và có fallback.",[e("page",59,"false positive làm UX tệ, user bypass system"),e("page",59,"Measure refuse rate, target ≤3%")],"Liên hệ over-filtering với hành vi người dùng.",multi=True),
        q("R24",R,"ambiguous","Judge score tốt là bao nhiêu?","",[],"Không rõ đang hỏi absolute rubric, pairwise win rate, Cohen kappa hay RAGAS metric.",answerability="ambiguous",ambiguity_reason="'Judge score' không xác định metric và mục đích.",interpretations=["absolute score 1–5","pairwise win rate","Cohen's kappa với human","RAGAS component score"],resolution="Cần chỉ rõ metric, judge và use case."),
        q("R25",R,"ambiguous","Guardrail nên block ở ngưỡng nào?","",[],"Tài liệu có nhiều classifier/layer và ngưỡng NLI khác nhau; câu hỏi không nêu loại guardrail hoặc risk profile.",answerability="ambiguous",difficulty="hard",ambiguity_reason="Thiếu classifier, miền rủi ro và hành động mong muốn.",interpretations=["NLI hallucination guardrail","topic validator","safety classifier"],resolution="Cần chỉ rõ guardrail, domain và chi phí false positive/negative."),
        q("R26",R,"insufficient evidence","Presidio đạt F1 bao nhiêu trên dữ liệu PII tiếng Việt của dự án?","",[],"Slide mô tả Presidio và regex tiếng Việt nhưng không báo cáo benchmark F1 trên dữ liệu dự án.",answerability="insufficient_evidence",difficulty="hard",nearest="Trang 47 nêu regex CCCD/phone/tax code và Presidio NER, không có F1."),
        q("R27",R,"contradictory evidence","Hai slide có thống nhất về ngưỡng NLI phải block hallucination không?","Chưa rõ: một slide nói dưới 0,5 là flag và dưới 0,3 là block, slide khác nói dưới 0,5 là block; cần xem đây là khác biệt policy/ngữ cảnh hay mâu thuẫn.",[e("page",39,"entailment_score < 0.5 →flag. < 0.3 →block."),e("page",58,"entailment_score < 0.5 →block")],"Natural inconsistency candidate; human phải phân loại direct/conditional/scope difference.",difficulty="hard",multi=True,contradiction=True,citation=True),
        q("R28",R,"citation-sensitive","Slide nêu những nghĩa vụ nào cho GDPR, EU AI Act và Vietnam PDPL?","Ứng viên trả lời: GDPR yêu cầu human override/audit cho quyết định; EU AI Act yêu cầu đánh giá phù hợp với hệ thống high-risk; Vietnam PDPL yêu cầu consent và DPIA khi chuyển dữ liệu cá nhân xuyên biên giới.",[e("page",65,"GDPR (EU) Article 22: no sole automated deci- sion; Article 13: explain logic"),e("page",65,"EU AI Act (Aug 2026 full) High-risk systems →conformity as- sessment"),e("page",65,"Vietnam PDPL (2025) Cross-border transfer của personal data cần consent + DPIA")],"Các tuyên bố pháp lý trong slide cần human xác minh và không được coi là tư vấn pháp lý.",difficulty="hard",multi=True,citation=True),

        q("S01",S,"exact-term factual","Social Engineering được định nghĩa như thế nào trong tiểu luận?","Là hành vi khai thác yếu tố con người để đánh lừa, lấy cắp thông tin hoặc tạo điều kiện cho tấn công mạng.",[e("paragraph",3,"Social Engineering (Kỹ thuật lừa đảo phi kỹ thuật) là hành vi khai thác yếu tố con người để đánh lừa, lấy cắp thông tin, hoặc tạo điều kiện cho tấn công mạng.")],"Định nghĩa trực tiếp."),
        q("S02",S,"exact-term factual","Trong tình huống CEO Fraud, kẻ tấn công giả danh ai và yêu cầu gì?","Giả danh giám đốc và yêu cầu kế toán chuyển tiền gấp.",[e("paragraph",9,"Vụ CEO Fraud: Hacker giả danh giám đốc gửi email yêu cầu kế toán chuyển tiền gấp → mất hàng trăm nghìn USD.")],"Chi tiết tình huống trực tiếp."),
        q("S03",S,"exact-term factual","Tiểu luận đề xuất những công nghệ hỗ trợ phòng chống nào?","2FA, bộ lọc thư rác/phishing và hệ thống phát hiện truy cập bất thường.",[e("paragraph",23,"Sử dụng xác thực 2 lớp (2FA), bộ lọc thư rác/phishing, hệ thống phát hiện truy cập bất thườg.")],"Danh sách biện pháp trực tiếp; giữ nguyên lỗi chính tả trong evidence.",citation=True),
        q("S04",S,"paraphrase factual","Vì sao lừa đảo phi kỹ thuật vẫn thường thành công dù không cần kỹ thuật cao?","Vì nó khai thác sự thiếu cảnh giác, lòng tin và thiếu kiến thức bảo mật của nạn nhân.",[e("paragraph",4,"Không cần kỹ thuật cao, Social Engineering thường thành công nhờ sự thiếu cảnh giác, lòng tin, và thiếu kiến thức bảo mật của nạn nhân.")],"Paraphrase nguyên nhân được nêu trực tiếp."),
        q("S05",S,"paraphrase factual","Chuỗi hậu quả bắt đầu từ email giả trong tình huống mô phỏng diễn ra thế nào?","Nhân viên nhập email, mật khẩu và OTP vào form giả; hacker đăng nhập, lấy dữ liệu khách hàng rồi tiếp tục lừa đảo.",[e("paragraph",12,"Nhân viên A nhận email từ địa chỉ giả it-support@xyz-corp.co, yêu cầu xác nhận tài khoản qua Google Form."),e("paragraph",13,"Nhân viên A nhập email, mật khẩu và OTP."),e("paragraph",14,"Hacker đăng nhập vào hệ thống, lấy dữ liệu khách hàng → tiếp tục lừa đảo.")],"Tổng hợp chuỗi sự kiện qua ba đoạn.",multi=True),
        q("S06",S,"multi-chunk","Tiểu luận kết hợp con người, quy trình và công nghệ để phòng chống Social Engineering ra sao?","Đào tạo/mô phỏng phishing nâng nhận thức; chính sách yêu cầu không cung cấp thông tin qua email và xác minh giao dịch; công nghệ gồm 2FA, lọc phishing và phát hiện truy cập bất thường.",[e("paragraph",17,"Tập huấn định kỳ, mô phỏng tấn công phishing."),e("paragraph",19,"Quy định không cung cấp thông tin qua email, xác minh 2 bước với các giao dịch quan trọng."),e("paragraph",23,"Sử dụng xác thực 2 lớp (2FA), bộ lọc thư rác/phishing, hệ thống phát hiện truy cập bất thườg.")],"Cần ba nhóm biện pháp ở nhiều đoạn.",multi=True,hierarchical=True),
        q("S07",S,"multi-chunk","Lab mô phỏng phishing ghi nhận hành vi gì và người dùng nên phản ứng thế nào?","Lab ghi nhận việc mở link/nhập dữ liệu giả; người dùng nên kiểm tra domain, không nhập dữ liệu đáng ngờ và báo phòng CNTT.",[e("paragraph",42,"Ghi nhận hành vi: người nhận có mở link không, có nhập thông tin không (dùng dữ liệu giả)."),e("paragraph",52,"Kiểm tra domain người gửi"),e("paragraph",53,"Không nhập dữ liệu nếu thấy nghi ngờ."),e("paragraph",54,"Báo ngay với phòng CNTT.")],"Kết hợp phần tấn công và phòng thủ của lab.",multi=True,hierarchical=True),
        q("S08",S,"cross-section","CEO Fraud trong phần tình huống liên hệ với dạng Whaling trong bảng như thế nào?","Cả hai đều mô tả giả mạo lãnh đạo/giám đốc để yêu cầu chuyển khoản gấp.",[e("paragraph",9,"Vụ CEO Fraud: Hacker giả danh giám đốc gửi email yêu cầu kế toán chuyển tiền gấp"),e("table",3,"3 | Whaling | Gửi email giả mạo giám đốc yêu cầu chuyển khoản gấp.")],"Liên kết ví dụ prose với taxonomy trong bảng.",multi=True,hierarchical=True),
        q("S09",S,"overview","Tiểu luận Social Engineering được tổ chức thành những phần chính nào?","Giới thiệu, các dạng tấn công, tình huống thực tế, biện pháp phòng chống, kết luận và phụ lục lab mô phỏng tấn công/phòng thủ.",[e("paragraph",2,"I. Giới thiệu chung"),e("paragraph",6,"II. Các dạng tấn công Social Engineering phổ biến"),e("paragraph",8,"III. Tình huống thực tế"),e("paragraph",15,"IV. Biện pháp phòng chống"),e("paragraph",25,"V. Kết luận"),e("paragraph",28,"PHỤ LỤC LAB MÔ PHỎNG: TẤN CÔNG & PHÒNG THỦ SOCIAL ENGINEERING")],"Câu hỏi cấu trúc toàn tài liệu.",multi=True,hierarchical=True),
        q("S10",S,"main-points","Kết luận của tiểu luận nhấn mạnh hai điểm nào?","Social Engineering nguy hiểm dù không đòi hỏi kỹ thuật cao; cần kết hợp đào tạo, quy trình và công nghệ vì con người là điểm yếu quan trọng.",[e("paragraph",26,"Social Engineering không đòi hỏi kỹ thuật cao nhưng cực kỳ nguy hiểm nếu người dùng mất cảnh giác."),e("paragraph",27,"Con người là điểm yếu quan trọng – cần kết hợp đào tạo, quy trình, công nghệ để đảm bảo an toàn thông tin.")],"Tóm tắt hai câu kết luận.",multi=True),
        q("S11",S,"compare","Phishing và Spear Phishing khác nhau thế nào theo bảng?","Phishing dùng email giả mạo chung có link form; Spear Phishing thêm tên và chức vụ cụ thể của người nhận để tăng thuyết phục.",[e("table",1,"1 | Phishing | Gửi email giả mạo có link Google Form yêu cầu nhập mật khẩu."),e("table",2,"2 | Spear Phishing | Gửi email có tên và chức vụ cụ thể người nhận, tăng tính thuyết phục.")],"Đối chiếu hai hàng taxonomy.",multi=True),
        q("S12",S,"how/why","Vì sao phòng chống Social Engineering không thể chỉ dựa vào công nghệ?","Vì tấn công khai thác con người; kết luận yêu cầu kết hợp đào tạo, quy trình và công nghệ.",[e("paragraph",3,"hành vi khai thác yếu tố con người"),e("paragraph",27,"cần kết hợp đào tạo, quy trình, công nghệ")],"Kết nối định nghĩa với kết luận.",multi=True),
        q("S13",S,"ambiguous","Địa chỉ email IT giả trong tài liệu là địa chỉ nào?","",[],"Tài liệu dùng hai địa chỉ khác nhau ở tình huống và lab.",answerability="ambiguous",ambiguity_reason="Có hai referent: địa chỉ trong tình huống thực tế và địa chỉ trong lab.",interpretations=["it-support@xyz-corp.co ở tình huống","it.support.xyzcorp@gmail.com ở lab"],resolution="Cần chỉ rõ phần tình huống hay phần lab."),
        q("S14",S,"insufficient evidence","Vụ CEO Fraud được nêu trong tiểu luận xảy ra vào năm nào và tại công ty nào?","",[],"Tiểu luận chỉ mô tả mô hình vụ việc và mức thiệt hại chung, không nêu năm hay công ty.",answerability="insufficient_evidence",nearest="Đoạn 9 nói mất hàng trăm nghìn USD nhưng không định danh vụ việc."),
    ]
    return rows


def _resolve(spec: tuple[str, int, str], doc_id: str, by_key: dict[tuple, dict]) -> dict:
    kind, number, needle = spec
    if kind == "page":
        locator = {"type": "pdf_page", "page": number}
    elif kind == "paragraph":
        locator = {"type": "docx_paragraph", "paragraph_index": number}
    elif kind == "table":
        locator = {"type": "docx_table_row", "table_index": 0, "row_index": number}
    else:
        raise ValueError(kind)
    unit = by_key[(doc_id, json.dumps(locator, ensure_ascii=False, sort_keys=True))]
    start = unit["text"].find(needle)
    if start < 0:
        raise ValueError(f"evidence not found for {doc_id} {locator}: {needle!r}")
    return {
        "locator": locator, "locator_label": unit["locator_label"],
        "start_offset": start, "end_offset": start + len(needle), "evidence_text": needle,
        "evidence_text_sha256": hashlib.sha256(needle.encode("utf-8")).hexdigest(),
    }


def _claims(answer: str) -> list[str]:
    return [part.strip().rstrip(".") for part in re.split(r";\s*|\.\s+(?=[A-ZÀ-Ỹ])", answer) if part.strip()]


def build(root: Path = ROOT) -> dict:
    root = Path(root)
    if any((root / name).read_text(encoding="utf-8").strip() for name in (
        "queries.jsonl", "canonical_evidence_spans.jsonl", "citation_annotations.jsonl"
    )):
        raise ValueError("refusing to regenerate candidates after human-approved records exist")
    docs = {r["doc_id"]: r for r in [json.loads(x) for x in (root / "documents.jsonl").read_text(encoding="utf-8").splitlines() if x.strip()]}
    units = source_units(root)
    by_key = {(u["doc_id"], json.dumps(u["locator"], ensure_ascii=False, sort_keys=True)): u for u in units}
    candidates, spans, claims = [], [], []
    for item in definitions():
        resolved = [_resolve(spec, item["doc"], by_key) for spec in item["evidence"]]
        doc = docs[item["doc"]]
        record = {
            "query_id": item["qid"], "document_id": item["doc"],
            "document_filename": doc["filename"], "split": doc["proposed_split"],
            "query": item["question"], "query_type": item["qtype"], "language": "vi",
            "difficulty": item["difficulty"], "answerability": item["answerability"],
            "candidate_answer": item["answer"], "candidate_evidence_spans": resolved,
            "candidate_rationale": item["rationale"], "requires_multi_chunk": item["multi"],
            "requires_hierarchical_context": item["hierarchical"],
            "citation_sensitive": item["citation"], "potential_contradiction": item["contradiction"],
            "ambiguity_reason": item["ambiguity_reason"],
            "acceptable_interpretations": item["interpretations"] or [],
            "resolution_evidence_needed": item["resolution"], "nearest_passage": item["nearest"],
            "candidate_generator": "LLM-assisted-source-grounded-v1",
            "annotation_status": "candidate",
        }
        candidates.append(record)
        for idx, span in enumerate(resolved, 1):
            spans.append({
                "span_id": f"{item['qid']}_candidate_s{idx:02d}", "query_id": item["qid"],
                "doc_id": item["doc"], **span, "annotation_status": "candidate",
            })
        if item["answerability"] == "answerable":
            support_ids = [f"{item['qid']}_candidate_s{i:02d}" for i in range(1, len(resolved) + 1)]
            for idx, claim in enumerate(_claims(item["answer"]), 1):
                claims.append({
                    "query_id": item["qid"], "answer_or_claim_id": f"{item['qid']}_c{idx:02d}",
                    "claim": claim, "supporting_span_ids": support_ids,
                    "support_label": "full", "support_label_is_proposal": True,
                    "annotation_status": "candidate",
                })
    contradiction_candidates = [
        {"pair_id":"NC01","query_id":"V16","span_a_id":"V16_candidate_s01","span_b_id":"V16_candidate_s02","proposed_label":"conditional difference","notes":"Metric mismatch warning versus equivalence under unit normalization.","human_rationale":"","annotation_status":"candidate"},
        {"pair_id":"NC02","query_id":"R27","span_a_id":"R27_candidate_s01","span_b_id":"R27_candidate_s02","proposed_label":"direct contradiction","notes":"Different block thresholds appear in two NLI policy slides; scope may differ.","human_rationale":"","annotation_status":"candidate"},
        {"pair_id":"NC03","query_id":"R14","span_a_id":"R14_candidate_s01","span_b_id":"R14_candidate_s02","proposed_label":"scope difference","notes":"End-of-day deliverable threshold versus CI gate design.","human_rationale":"","annotation_status":"candidate"},
    ]
    artifact_candidates = [
        {"document_id":"cv1_vector_feature_store","evaluation_scope":["summary","mind_map"],"important_concepts":["embeddings","vector database","ANN and hybrid retrieval","RAG","GraphRAG","feature store"],"important_sections":["embedding models","similarity metrics","hybrid search","RAG pipeline","GraphRAG","feature stores"],"expected_relations":["embedding model -> vector index","BM25 + vector -> RRF","documents -> knowledge graph","feature store -> online/offline store"],"critical_facts":["Changing embedding model requires reindexing","Index and query must use the same model version"],"critical_exclusions":["Do not invent details hidden in slide graphics without visual confirmation"],"invalid_or_redundant_relations":[],"annotator":{"count":0,"agreement":"not applicable / not measured"},"annotation_status":"candidate"},
        {"document_id":"cv1_ragas_guardrails","evaluation_scope":["summary","mind_map"],"important_concepts":["evaluation levels","RAGAS metrics","LLM-as-Judge biases","hallucination detection","guardrail axes","defense-in-depth","CI/CD evaluation"],"important_sections":["Foundations","RAGAS","LLM-as-Judge","Hallucination Detection","Guardrails","Production Patterns"],"expected_relations":["Faithfulness -> hallucination","Context Precision -> ranking","judge bias -> mitigation","guardrail axis -> layered control","production failure -> regression case"],"critical_facts":["Human labels are still needed to calibrate LLM judges","Offline and online evaluation are complementary"],"critical_exclusions":["Do not present slide benchmark targets as universal guarantees","Do not treat legal slide content as legal advice"],"invalid_or_redundant_relations":[],"annotator":{"count":0,"agreement":"not applicable / not measured"},"annotation_status":"candidate"},
        {"document_id":"cv1_social_engineering","evaluation_scope":["summary","mind_map"],"important_concepts":["human-factor exploitation","attack types","CEO fraud","phishing simulation","awareness training","policy","2FA"],"important_sections":["Introduction","Attack types","Scenarios","Countermeasures","Conclusion","Simulation lab"],"expected_relations":["phishing -> credential theft","CEO fraud -> urgent transfer request","training + process + technology -> defense","suspicious email -> verify domain/report IT"],"critical_facts":["Social Engineering exploits human factors","Defense combines training, process, and technology"],"critical_exclusions":["Do not assert a date or named company for the CEO Fraud example"],"invalid_or_redundant_relations":[],"annotator":{"count":0,"agreement":"not applicable / not measured"},"annotation_status":"candidate"},
    ]
    outputs = {
        "source_units.jsonl": units, "query_candidates.jsonl": candidates,
        "candidate_canonical_evidence_spans.jsonl": spans,
        "candidate_citation_annotations.jsonl": claims,
        "candidate_contradictions.jsonl": contradiction_candidates,
        "candidate_artifact_annotations.jsonl": artifact_candidates,
    }
    for filename, rows in outputs.items():
        (root / filename).write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")
    from .review import export_review
    export_review(root)
    summary = {
        "generated_at": datetime.now(timezone.utc).isoformat(), "annotation_status": "candidate",
        "warning": "LLM-assisted annotation aids only; none are ground truth until explicit human import.",
        "candidate_query_count": len(candidates),
        "query_type_distribution": dict(sorted(Counter(c["query_type"] for c in candidates).items())),
        "document_distribution": dict(sorted(Counter(c["document_id"] for c in candidates).items())),
        "split_distribution": dict(sorted(Counter(c["split"] for c in candidates).items())),
        "answerability_distribution": dict(sorted(Counter(c["answerability"] for c in candidates).items())),
        "candidate_span_count": len(spans), "candidate_claim_count": len(claims),
        "candidate_contradiction_pair_count": len(contradiction_candidates),
        "candidate_artifact_document_count": len(artifact_candidates),
        "annotator_count": 0, "inter_annotator_agreement": "not applicable / not measured",
    }
    (root / "candidate_package_manifest.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    hash_files = [
        root / "documents.jsonl", root / "query_candidates.jsonl",
        root / "candidate_canonical_evidence_spans.jsonl",
        root / "candidate_citation_annotations.jsonl", root / "candidate_contradictions.jsonl",
        root / "candidate_artifact_annotations.jsonl", root / "source_units.jsonl",
        root / "human_review_queries.csv", root / "human_review_citations.csv",
        root / "human_review_contradictions.csv", root / "human_review_artifacts.csv",
    ]
    file_hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in hash_files}
    package_hash = hashlib.sha256(
        "".join(f"{name}:{value}\n" for name, value in sorted(file_hashes.items())).encode()
    ).hexdigest()
    manifest_path = root / "dataset_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest.update({
        "dataset_version": "corpus_v1-candidate.2", "corpus_version": "corpus_v1-candidate.2",
        "status": "human_validation_package_ready", "query_count": 0,
        "candidate_query_count": len(candidates),
        "candidate_query_type_distribution": summary["query_type_distribution"],
        "candidate_answerability_distribution": summary["answerability_distribution"],
        "candidate_split_distribution": summary["split_distribution"],
        "proposed_split": {
            "development": ["cv1_ragas_guardrails"],
            "test": ["cv1_vector_feature_store", "cv1_social_engineering"],
            "ratio": "1/3 development; 2/3 test at document level",
            "limitation": "Only three real documents exist; AI/RAG is represented in both splits, but cybersecurity occurs only in test.",
        },
        "annotation_status": {
            "documents": "candidate_human_confirmation_pending",
            "queries": "60_candidates_0_human_validated",
            "qrels": "not_started",
            "gold_answers": "candidate_aids_only",
            "canonical_evidence_spans": "92_candidates_0_human_validated",
            "citations": "83_candidate_claims_0_human_validated",
            "contradictions": "3_candidate_pairs_0_human_validated",
            "artifacts": "3_candidate_documents_0_human_validated",
        },
        "split_status": "proposed_not_frozen", "hash_status": "mutable_candidate_snapshot",
        "candidate_package_hash": package_hash, "candidate_package_file_hashes": file_hashes,
        "dataset_hash": None, "split_hash": None, "annotation_hash": None,
        "freeze_blocked": True,
        "freeze_blockers": [
            "All 60 query records and their proposed answers/evidence remain candidate-only.",
            "No human has approved corpus membership or the proposed document split.",
            "The eligible corpus contains only three genuine documents; no additional real documents were found.",
            "R0/R1/R2 must not be built until corpus selection and source hashes are human-approved.",
        ],
    })
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    summary["candidate_package_hash"] = package_hash
    (root / "candidate_package_manifest.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return summary


def main(argv=None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset", type=Path, default=ROOT)
    args = parser.parse_args(argv)
    print(json.dumps(build(args.dataset), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
