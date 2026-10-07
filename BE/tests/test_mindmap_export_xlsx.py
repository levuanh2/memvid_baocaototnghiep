import io

from openpyxl import load_workbook

from services.mindmap.export.scope import resolve_export_scope
from services.mindmap.export.tree import build_export_tree
from services.mindmap.export.xlsx_serializer import serialize_xlsx


def _nodes():
    return [
        {"id": "r", "parent": None, "kind": "root", "title": "Bản đồ tư duy", "order": 0},
        {"id": "c1", "parent": "r", "kind": "section", "title": "Kiến trúc hệ thống", "order": 0, "chunk_refs": ["doc1#p3"]},
        {"id": "c1a", "parent": "c1", "kind": "idea", "title": "Chi tiết A", "order": 0},
        {"id": "c2", "parent": "r", "kind": "section", "title": "Bảo mật", "order": 1},
    ]


def _tree(nodes=None, relations=None):
    nodes = nodes or _nodes()
    scope = resolve_export_scope(nodes, scope_type="full")
    return build_export_tree(nodes, relations or [], scope["root_ids"], scope["included_ids"], map_id="m1", title="Bản đồ tư duy")


def test_xlsx_has_summary_and_nodes_sheets_with_correct_row_count():
    wb = load_workbook(io.BytesIO(serialize_xlsx(_tree())))
    assert "Summary" in wb.sheetnames
    assert "Nodes" in wb.sheetnames
    nodes_ws = wb["Nodes"]
    header = [c.value for c in nodes_ws[1]]
    assert header == ["node_id", "parent_id", "branch_path", "depth", "order", "topic", "is_context"]
    # header + 4 nodes (r, c1, c1a, c2)
    assert nodes_ws.max_row == 5


def test_xlsx_node_ids_parent_relationships_branch_path_depth_order_and_vietnamese_text():
    wb = load_workbook(io.BytesIO(serialize_xlsx(_tree())))
    rows = {row[0].value: row for row in wb["Nodes"].iter_rows(min_row=2)}
    c1a = rows["c1a"]
    assert c1a[1].value == "c1"  # parent_id
    assert c1a[2].value == "Bản đồ tư duy / Kiến trúc hệ thống / Chi tiết A"  # branch_path
    assert c1a[3].value == 2  # depth
    assert isinstance(c1a[4].value, int)  # order
    assert c1a[5].value == "Chi tiết A"  # topic (Vietnamese diacritics intact)


def test_xlsx_citations_sheet_only_appears_when_requested_and_present():
    with_cit = load_workbook(io.BytesIO(serialize_xlsx(_tree(), include_citations=True)))
    assert "Citations" in with_cit.sheetnames
    rows = list(with_cit["Citations"].iter_rows(min_row=2, values_only=True))
    assert rows == [("c1", "Kiến trúc hệ thống", "doc1#p3")]

    without_cit = load_workbook(io.BytesIO(serialize_xlsx(_tree(), include_citations=False)))
    assert "Citations" not in without_cit.sheetnames

    no_citations_tree = _tree([{**n, "chunk_refs": []} for n in _nodes()])
    no_sheet = load_workbook(io.BytesIO(serialize_xlsx(no_citations_tree, include_citations=True)))
    assert "Citations" not in no_sheet.sheetnames  # nothing to show -> sheet omitted, not empty


def test_xlsx_relations_sheet_only_appears_when_requested_and_present():
    relations = [{"source": "c1a", "target": "c2", "type": "relates_to", "label": "liên quan"}]
    tree_with_rel = _tree(relations=relations)
    wb = load_workbook(io.BytesIO(serialize_xlsx(tree_with_rel, include_relations=True)))
    assert "Relations" in wb.sheetnames
    row = list(wb["Relations"].iter_rows(min_row=2, values_only=True))[0]
    assert row == ("c1a", "Chi tiết A", "c2", "Bảo mật", "relates_to", "liên quan")

    wb_excluded = load_workbook(io.BytesIO(serialize_xlsx(tree_with_rel, include_relations=False)))
    assert "Relations" not in wb_excluded.sheetnames

    no_rel_tree = _tree()
    wb_none = load_workbook(io.BytesIO(serialize_xlsx(no_rel_tree, include_relations=True)))
    assert "Relations" not in wb_none.sheetnames


def test_xlsx_notes_and_source_names_columns_are_opt_in():
    nodes = _nodes()
    nodes[1]["chunk_refs"] = ["doc1#p3", "doc2#p7"]  # c1 gets citations
    tree = _tree(nodes)
    default_wb = load_workbook(io.BytesIO(serialize_xlsx(tree)))
    header_default = [c.value for c in default_wb["Nodes"][1]]
    assert "note" not in header_default
    assert "source_names" not in header_default

    full_wb = load_workbook(io.BytesIO(serialize_xlsx(tree, content={"notes": True, "sourceNames": True})))
    header_full = [c.value for c in full_wb["Nodes"][1]]
    assert "note" in header_full
    assert "source_names" in header_full
    rows = {row[0].value: row for row in full_wb["Nodes"].iter_rows(min_row=2)}
    source_names_col = header_full.index("source_names")
    assert "doc1" in rows["c1"][source_names_col].value
    assert "doc2" in rows["c1"][source_names_col].value


def test_xlsx_font_selection_applied_to_header_and_body():
    wb = load_workbook(io.BytesIO(serialize_xlsx(_tree(), font="serif")))
    ws = wb["Nodes"]
    assert ws["A1"].font.name == "Times New Roman"
    assert ws["A2"].font.name == "Times New Roman"


def test_xlsx_header_style_monochrome_and_custom_palette_fill_the_header_row():
    keep_wb = load_workbook(io.BytesIO(serialize_xlsx(_tree(), header_style_mode="keep")))
    mono_wb = load_workbook(io.BytesIO(serialize_xlsx(_tree(), header_style_mode="monochrome")))
    palette_wb = load_workbook(io.BytesIO(serialize_xlsx(_tree(), header_style_mode="customPalette")))
    assert keep_wb["Nodes"]["A1"].fill.fgColor.rgb in (None, "00000000")
    assert mono_wb["Nodes"]["A1"].fill.fgColor.rgb == "002B2620" or mono_wb["Nodes"]["A1"].fill.fgColor.rgb == "FF2B2620"
    assert palette_wb["Nodes"]["A1"].fill.fgColor.rgb != mono_wb["Nodes"]["A1"].fill.fgColor.rgb


def test_xlsx_header_row_frozen_and_autofilter_enabled():
    wb = load_workbook(io.BytesIO(serialize_xlsx(_tree())))
    ws = wb["Nodes"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None


def test_xlsx_column_widths_are_sensible_not_default():
    wb = load_workbook(io.BytesIO(serialize_xlsx(_tree())))
    ws = wb["Nodes"]
    # "branch_path" column holds long strings — its width must reflect that,
    # not openpyxl's bare default (~8.43).
    branch_path_col = [c.value for c in ws[1]].index("branch_path")
    from openpyxl.utils import get_column_letter
    letter = get_column_letter(branch_path_col + 1)
    assert ws.column_dimensions[letter].width > 15


def test_xlsx_excludes_scope_outside_selection():
    nodes = _nodes()
    scope = resolve_export_scope(nodes, scope_type="current_branch", selected_node_id="c1")
    tree = build_export_tree(nodes, [], scope["root_ids"], scope["included_ids"], title="T")
    wb = load_workbook(io.BytesIO(serialize_xlsx(tree)))
    topics = [row[0] for row in wb["Nodes"].iter_rows(min_row=2, min_col=6, max_col=6, values_only=True)]
    assert "Bảo mật" not in topics
    assert "Chi tiết A" in topics
