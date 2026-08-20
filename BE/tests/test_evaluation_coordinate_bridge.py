from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from evaluation.evidence import (
    bridge_span_to_canonical,
    canonical_record_for_document,
    collapse_whitespace_with_offsets,
)


def _fixture():
    canonical_text = "Alpha\n beta  gamma"
    source_text = "Alpha beta gamma"
    locator = {"type": "pdf_page", "page": 1}
    source_unit = {
        "doc_id": "doc-a",
        "locator": locator,
        "text": source_text,
        "text_sha256": hashlib.sha256(source_text.encode()).hexdigest(),
    }
    record = {
        "doc_id": "doc-a",
        "canonical_text": canonical_text,
        "canonical_text_sha256": hashlib.sha256(canonical_text.encode()).hexdigest(),
        "units": [{"locator": locator, "char_start": 0, "char_end": len(canonical_text)}],
    }
    return source_unit, record


def _span(text: str, start: int, end: int, locator: dict | None = None) -> dict:
    return {
        "span_id": "s1",
        "doc_id": "doc-a",
        "locator": locator or {"type": "pdf_page", "page": 1},
        "start_offset": start,
        "end_offset": end,
        "evidence_text": text,
        "evidence_text_sha256": hashlib.sha256(text.encode()).hexdigest(),
    }


def test_coordinate_bridge_exact_mapping():
    unit, record = _fixture()
    bridged = bridge_span_to_canonical(_span("beta", 6, 10), doc_id="doc-a", source_unit=unit, canonical_record=record)
    assert bridged["locator"]["char_start"] == 7
    assert bridged["locator"]["char_end"] == 11
    assert bridged["source_locator"] == {"type": "pdf_page", "page": 1}
    assert bridged["coordinate_bridge"]["mapping"] == "exact_source_unit_to_exact_canonical_unit"


def test_coordinate_bridge_mapping_at_both_boundaries():
    unit, record = _fixture()
    bridged = bridge_span_to_canonical(
        _span(unit["text"], 0, len(unit["text"])),
        doc_id="doc-a", source_unit=unit, canonical_record=record,
    )
    assert bridged["locator"]["char_start"] == 0
    assert bridged["locator"]["char_end"] == len(record["canonical_text"])


def test_coordinate_bridge_multiple_spans_same_document():
    unit, record = _fixture()
    alpha = bridge_span_to_canonical(_span("Alpha", 0, 5), doc_id="doc-a", source_unit=unit, canonical_record=record)
    gamma = bridge_span_to_canonical(_span("gamma", 11, 16), doc_id="doc-a", source_unit=unit, canonical_record=record)
    assert (alpha["locator"]["char_start"], alpha["locator"]["char_end"]) == (0, 5)
    assert record["canonical_text"][gamma["locator"]["char_start"]:gamma["locator"]["char_end"]] == "gamma"


def test_coordinate_bridge_rejects_invalid_document_id():
    unit, record = _fixture()
    with pytest.raises(ValueError, match="document identity mismatch"):
        bridge_span_to_canonical(_span("beta", 6, 10), doc_id="doc-b", source_unit=unit, canonical_record=record)


def test_coordinate_bridge_rejects_out_of_range_offset():
    unit, record = _fixture()
    with pytest.raises(ValueError, match="offsets out of range"):
        bridge_span_to_canonical(_span("gamma", 11, 99), doc_id="doc-a", source_unit=unit, canonical_record=record)


def test_coordinate_bridge_rejects_unmappable_span():
    unit, record = _fixture()
    missing = _span("beta", 6, 10, {"type": "pdf_page", "page": 2})
    unit = {**unit, "locator": missing["locator"]}
    with pytest.raises(ValueError, match="unmappable source locator"):
        bridge_span_to_canonical(missing, doc_id="doc-a", source_unit=unit, canonical_record=record)


def test_coordinate_bridge_is_deterministic_on_repeated_mapping():
    unit, record = _fixture()
    span = _span("beta", 6, 10)
    first = bridge_span_to_canonical(span, doc_id="doc-a", source_unit=unit, canonical_record=record)
    second = bridge_span_to_canonical(span, doc_id="doc-a", source_unit=unit, canonical_record=record)
    assert first == second


def test_coordinate_bridge_preserves_source_text_while_mapping_nfc_offsets():
    source_text = "e\u0301 fact"
    canonical_text = "é fact"
    locator = {"type": "pdf_page", "page": 1}
    unit = {
        "doc_id": "doc-a", "locator": locator, "text": source_text,
        "text_sha256": hashlib.sha256(source_text.encode()).hexdigest(),
    }
    record = {
        "doc_id": "doc-a", "canonical_text": canonical_text,
        "canonical_text_sha256": hashlib.sha256(canonical_text.encode()).hexdigest(),
        "units": [{"locator": locator, "char_start": 0, "char_end": len(canonical_text)}],
    }
    reviewed = _span("e\u0301", 0, 2)
    bridged = bridge_span_to_canonical(reviewed, doc_id="doc-a", source_unit=unit, canonical_record=record)
    assert bridged["source_evidence_text"] == "e\u0301"
    assert bridged["evidence_text"] == "é"
    assert (bridged["locator"]["char_start"], bridged["locator"]["char_end"]) == (0, 1)


def test_docx_canonical_record_includes_reviewable_table_rows(tmp_path: Path):
    from docx import Document

    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Paragraph evidence")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cell A"
    table.cell(0, 1).text = "Cell B"
    doc.save(source)
    record = canonical_record_for_document({
        "doc_id": "docx-a", "source": str(source), "format": "docx",
        "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
    })
    locators = [unit["locator"] for unit in record["units"]]
    assert {"type": "docx_paragraph", "paragraph_index": 0} in locators
    assert {"type": "docx_table_row", "table_index": 0, "row_index": 0} in locators
    assert "Cell A | Cell B" in record["canonical_text"]


def test_whitespace_offset_map_rejects_non_nfc_canonical_text():
    with pytest.raises(ValueError, match="must be NFC"):
        collapse_whitespace_with_offsets("e\u0301")
