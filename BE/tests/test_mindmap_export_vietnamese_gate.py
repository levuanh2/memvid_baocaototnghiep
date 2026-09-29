"""Section 7 (final hardening round): Vietnamese text must render correctly
in every format's EXTRACTED/RENDERED output, not just survive as an input
string. Uses the exact representative character set the round specified:
ă â ê ô ơ ư đ Á Ế Ỗ Ờ Ữ — embedded in real node topics/notes, then the
generated file is reopened with the SAME library that wrote it and the
text is read back out.

PNG/JPEG/SVG (client-side, browser-rendered) are covered separately by
FE/e2e-fixture/export-studio.spec.js's real-Chromium assertions (SVG's
"Kiến trúc hệ thống" etc. already exercises Vietnamese diacritics via a
real XML parse of the downloaded file) — no backend rendering surface
exists for those formats to test here.
"""
from __future__ import annotations

import io

import docx
import fitz
from openpyxl import load_workbook

from services.mindmap.export.scope import resolve_export_scope
from services.mindmap.export.tree import build_export_tree
from services.mindmap.export.docx_serializer import serialize_docx
from services.mindmap.export.xlsx_serializer import serialize_xlsx
from services.mindmap.export.pdf_serializer import serialize_pdf

# The exact representative string this round specified.
VN_TEXT = "ă â ê ô ơ ư đ Á Ế Ỗ Ờ Ữ"


def _tree_with_vn_text():
    nodes = [
        {"id": "root", "parent": None, "kind": "root", "title": "Bản đồ tư duy", "order": 0},
        {"id": "c1", "parent": "root", "kind": "section", "title": VN_TEXT, "order": 0, "note": VN_TEXT},
    ]
    scope = resolve_export_scope(nodes, scope_type="full")
    return build_export_tree(nodes, [], scope["root_ids"], scope["included_ids"], title=VN_TEXT)


def test_docx_extracted_text_contains_the_exact_vietnamese_string():
    tree = _tree_with_vn_text()
    doc = docx.Document(io.BytesIO(serialize_docx(tree, content={"notes": True})))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert VN_TEXT in full_text


def test_xlsx_extracted_cell_values_contain_the_exact_vietnamese_string():
    tree = _tree_with_vn_text()
    wb = load_workbook(io.BytesIO(serialize_xlsx(tree)))
    values = [c.value for row in wb["Nodes"].iter_rows() for c in row]
    assert VN_TEXT in values


def test_pdf_extracted_text_contains_the_exact_vietnamese_string_sans_and_serif():
    tree = _tree_with_vn_text()
    for font in ("sans", "serif"):
        file_bytes = serialize_pdf(tree, mode="outline", font=font, content={"notes": True})
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        full_text = "".join(page.get_text() for page in doc)
        assert VN_TEXT in full_text, font
        doc.close()


def test_pdf_extracted_text_survives_every_branch_color_mode():
    """Color mode changes the drawing operator's color argument, never the
    glyph/text-run itself — confirms that isn't silently corrupting text."""
    tree = _tree_with_vn_text()
    for mode in ("keep", "monochrome", "customPalette"):
        file_bytes = serialize_pdf(tree, mode="outline", branch_color_mode=mode)
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        full_text = "".join(page.get_text() for page in doc)
        assert "Á Ế Ỗ Ờ Ữ" in full_text, mode
        doc.close()
