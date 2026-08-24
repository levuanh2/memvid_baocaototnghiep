"""Định dạng tài liệu upload: một nguồn sự thật, và các bộ đọc thật sự đọc được.

Ba thứ được khoá ở đây:
  1. Chuỗi `accept` của FE khớp `formats.SUPPORTED_EXTENSIONS` — trước đây FE,
     document_loader và markdown_convert giữ ba danh sách lệch nhau nên
     `.csv`/`.json` bị giấu khỏi người dùng còn `.html` thì chết lặng.
  2. Mỗi bộ đọc mới trả về nội dung thật, kiểm bằng file sinh tại chỗ chứ không
     phải fixture cam kết sẵn.
  3. `/upload-file` từ chối đuôi lạ bằng 415 và KHÔNG ghi gì xuống đĩa.
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import pytest

from app.domains.ingest import formats
from app.domains.ingest.document_loader import load_document

# Neo theo vị trí file test, KHÔNG theo CWD — pytest chạy từ BE/ nên đường dẫn
# tương đối trần sẽ trỏ nhầm (đúng cái bẫy làm 17 test evaluation chết lặng).
REPO_ROOT = Path(__file__).resolve().parents[2]
DOCUMENT_LIST_JSX = REPO_ROOT / "FE" / "src" / "pages" / "study" / "DocumentList.jsx"


# ── Nguồn sự thật ────────────────────────────────────────────────────────────

def test_fe_accept_khop_backend():
    source = DOCUMENT_LIST_JSX.read_text(encoding="utf-8")
    match = re.search(r'accept="([^"]+)"', source)
    assert match, "khong tim thay thuoc tinh accept trong DocumentList.jsx"
    fe_exts = set(match.group(1).split(","))
    assert fe_exts == formats.SUPPORTED_EXTENSIONS, (
        "FE va BE lech nhau: "
        f"FE thua {sorted(fe_exts - formats.SUPPORTED_EXTENSIONS)}, "
        f"FE thieu {sorted(formats.SUPPORTED_EXTENSIONS - fe_exts)}"
    )


def test_is_supported_khong_phan_biet_hoa_thuong():
    assert formats.is_supported("Bai giang.PPTX")
    assert formats.is_supported("giao-trinh.epub")
    assert not formats.is_supported("virus.exe")
    assert not formats.is_supported("khong-co-duoi")
    assert not formats.is_supported("")


def test_moi_dinh_dang_thuoc_dung_mot_nhom():
    """Mỗi đuôi chỉ nằm trong một nhóm — trùng nhóm nghĩa là dispatch nhập nhằng."""
    groups = [formats.TEXT, formats.PYMUPDF, formats.WORD, formats.PPTX,
              formats.XLSX, formats.OPENDOCUMENT, formats.WEB, formats.DATA,
              formats.RICH_TEXT, formats.IMAGE]
    tong = sum(len(g) for g in groups)
    assert tong == len(formats.SUPPORTED_EXTENSIONS), "co duoi nam o hai nhom"


# ── Bộ đọc ───────────────────────────────────────────────────────────────────

def test_html_doc_duoc(tmp_path: Path):
    """Hồi quy: bản cũ gọi UnstructuredHTMLLoader (không có trong requirements),
    lỗi import bị `except Exception: pass` nuốt nên mọi file .html trả rỗng và
    ingest chết với "Cannot read file content"."""
    f = tmp_path / "bai.html"
    f.write_text("<h1>Dao ham</h1><p>Quy tac chuoi</p>", encoding="utf-8")
    docs = load_document(str(f))
    assert docs, "html phai doc duoc"
    assert "Dao ham" in docs[0].page_content


def test_pptx_moi_slide_mot_document(tmp_path: Path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    deck = pptx.Presentation()
    s1 = deck.slides.add_slide(deck.slide_layouts[1])
    s1.shapes.title.text = "Gioi han"
    s1.notes_slide.notes_text_frame.text = "Ghi chu giang bai"
    s2 = deck.slides.add_slide(deck.slide_layouts[5])
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Ham"
    table.cell(1, 0).text = "sin x"
    f = tmp_path / "slide.pptx"
    deck.save(str(f))

    docs = load_document(str(f))
    assert len(docs) == 2
    assert docs[0].metadata["slide"] == 1
    assert "Ghi chu giang bai" in docs[0].page_content, "ghi chu slide phai duoc giu"
    assert "sin x" in docs[1].page_content, "o bang phai duoc doc"


def test_xlsx_moi_sheet_mot_bang_markdown(tmp_path: Path):
    openpyxl = pytest.importorskip("openpyxl")

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Diem"
    for row in [["Ten", "Diem"], ["An", 9]]:
        sheet.append(row)
    book.create_sheet("Trong")   # sheet rong -> phai bi bo qua
    f = tmp_path / "bang.xlsx"
    book.save(str(f))

    docs = load_document(str(f))
    assert len(docs) == 1, "sheet rong khong duoc tao Document"
    assert docs[0].metadata["sheet"] == "Diem"
    assert "| Ten | Diem |" in docs[0].page_content


def test_rtf_doc_duoc(tmp_path: Path):
    pytest.importorskip("striprtf")
    f = tmp_path / "ghi-chu.rtf"
    f.write_text(chr(123) + chr(92) + "rtf1" + chr(92) + "ansi Tich phan tung phan."
                 + chr(125), encoding="utf-8")
    docs = load_document(str(f))
    assert docs and "Tich phan" in docs[0].page_content


def test_epub_doc_bang_pymupdf(tmp_path: Path):
    """PyMuPDF mở EPUB native — không cần thư viện ebook riêng."""
    f = tmp_path / "sach.epub"
    with zipfile.ZipFile(f, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container version="1.0" '
                   'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                   '<rootfile full-path="OEBPS/content.opf" '
                   'media-type="application/oebps-package+xml"/></rootfiles></container>')
        z.writestr("OEBPS/content.opf",
                   '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" '
                   'version="2.0" unique-identifier="i"><metadata '
                   'xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Sach</dc:title>'
                   '<dc:identifier id="i">x1</dc:identifier><dc:language>vi</dc:language>'
                   '</metadata><manifest><item id="c1" href="c1.xhtml" '
                   'media-type="application/xhtml+xml"/></manifest><spine>'
                   '<itemref idref="c1"/></spine></package>')
        z.writestr("OEBPS/c1.xhtml",
                   '<?xml version="1.0"?><!DOCTYPE html><html '
                   'xmlns="http://www.w3.org/1999/xhtml"><head><title>C1</title></head>'
                   '<body><h1>Chuong mot</h1><p>Noi dung sach.</p></body></html>')

    docs = load_document(str(f))
    assert docs and "Chuong mot" in docs[0].page_content


def test_duoi_la_tra_rong_chu_khong_no(tmp_path: Path):
    f = tmp_path / "rac.exe"
    f.write_bytes(b"MZ binary")
    assert load_document(str(f)) == []


# ── Cổng chặn ở route ────────────────────────────────────────────────────────

def test_upload_tu_choi_duoi_la_va_khong_ghi_dia(client, monkeypatch, tmp_path: Path):
    import app.main as main

    monkeypatch.setattr(main, "_auth_protect_enabled", lambda: False)
    monkeypatch.setattr(main, "_current_user_id", lambda: None)
    monkeypatch.setattr(main, "INPUT_DIR", str(tmp_path))

    resp = client.post(
        "/upload-file",
        data={"file": (io.BytesIO(b"MZ binary"), "malware.exe")},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 415
    payload = resp.get_json()
    assert ".pptx" in payload["supported_extensions"]
    assert list(tmp_path.iterdir()) == [], "file bi tu choi khong duoc ghi xuong dia"


def test_co_gioi_han_dung_luong():
    import app.main as main

    limit = main.app.config.get("MAX_CONTENT_LENGTH")
    assert limit and limit > 0, "thieu MAX_CONTENT_LENGTH -> upload khong gioi han"
    assert limit == main.MAX_UPLOAD_MB * 1024 * 1024
