"""PR B (side-agent note follow-up): an ancestor kept only as orientation
context ("is_context") must be visually distinguishable from a node the user
actually selected in every document format — otherwise "Root" and "A" render
as ordinary headings with no content of their own under a current_branch/
selected_branches export, reading as broken rather than intentional. No
italic variant of the bundled DejaVu fonts exists for PDF (assets/fonts has
only the regular weight), so PDF uses muted gray text for context lines
instead of italic; DOCX and XLSX use italic.
"""
import io

import docx
import fitz
from openpyxl import load_workbook

from services.mindmap.export.scope import resolve_export_scope
from services.mindmap.export.tree import build_export_tree
from services.mindmap.export.docx_serializer import serialize_docx
from services.mindmap.export.pdf_serializer import serialize_pdf, CONTEXT_TEXT_RGB
from services.mindmap.export.xlsx_serializer import serialize_xlsx


def _nodes():
    return [
        {"id": "r", "parent": None, "kind": "root", "title": "Bản đồ tư duy", "order": 0},
        {"id": "c1", "parent": "r", "kind": "section", "title": "Kiến trúc hệ thống", "order": 0},
        {"id": "c1a", "parent": "c1", "kind": "idea", "title": "Chi tiết A", "order": 0},
        {"id": "c2", "parent": "r", "kind": "section", "title": "Bảo mật", "order": 1},
    ]


def _branch_tree():
    """current_branch on "c1a" -> included={c1a}, context={r, c1} — exactly the
    shape the side-agent note described."""
    nodes = _nodes()
    scope = resolve_export_scope(nodes, scope_type="current_branch", selected_node_id="c1a")
    return build_export_tree(
        nodes, [], scope["root_ids"], scope["included_ids"],
        map_id="m1", title="Bản đồ tư duy", context_ids=scope["context_ids"],
    )


def test_docx_context_heading_is_italic_selected_heading_is_not():
    doc = docx.Document(io.BytesIO(serialize_docx(_branch_tree())))
    italic_by_text = {}
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            italic_by_text[p.text] = any(r.italic for r in p.runs)
    assert italic_by_text["Bản đồ tư duy"] is True  # root — context only
    assert italic_by_text["Kiến trúc hệ thống"] is True  # c1 — context only
    assert italic_by_text["Chi tiết A"] is False  # c1a — the actual selection


def test_pdf_context_lines_are_gray_selected_line_is_not():
    pdf_bytes = serialize_pdf(_branch_tree(), mode="outline", content={})
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    spans_by_text = {}
    for page in doc:
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    spans_by_text[span["text"]] = span["color"]
    doc.close()

    def rgb01(packed_int):
        return ((packed_int >> 16) & 255) / 255, ((packed_int >> 8) & 255) / 255, (packed_int & 255) / 255

    def close(a, b, tol=0.02):
        return all(abs(x - y) < tol for x, y in zip(a, b))

    assert close(rgb01(spans_by_text["Bản đồ tư duy"]), CONTEXT_TEXT_RGB)
    assert close(rgb01(spans_by_text["Kiến trúc hệ thống"]), CONTEXT_TEXT_RGB)
    assert not close(rgb01(spans_by_text["Chi tiết A"]), CONTEXT_TEXT_RGB)


def test_xlsx_nodes_sheet_has_is_context_column_and_italic_context_topic_cell():
    wb = load_workbook(io.BytesIO(serialize_xlsx(_branch_tree())))
    ws = wb["Nodes"]
    header = [c.value for c in ws[1]]
    is_context_col = header.index("is_context")
    topic_col = header.index("topic")
    by_id = {row[0].value: row for row in ws.iter_rows(min_row=2)}
    assert by_id["r"][is_context_col].value is True
    assert by_id["c1"][is_context_col].value is True
    assert by_id["c1a"][is_context_col].value is False
    assert by_id["r"][topic_col].font.italic is True
    assert by_id["c1"][topic_col].font.italic is True
    assert not by_id["c1a"][topic_col].font.italic
